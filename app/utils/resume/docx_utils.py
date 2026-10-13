import io
import zipfile
import xml.etree.ElementTree as ET
from typing import Dict, List, Optional
import tempfile
import logging

from pdf2docx import Converter
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import os

from .pdf_utils import make_simple_pdf_from_text

logger = logging.getLogger(__name__)

def is_docx_bytes(data: bytes) -> bool:
    return isinstance(data, (bytes, bytearray)) and data[:4] == b"PK\x03\x04"


def _get_relationships(zf: zipfile.ZipFile, filename: str) -> Dict[str, str]:
    """Parses .rels file to map rId to Target (URL)."""
    rels = {}
    if filename not in zf.namelist():
        return rels
    try:
        with zf.open(filename) as f:
            tree = ET.fromstring(f.read())
        ns = {"r": "http://schemas.openxmlformats.org/package/2006/relationships"}
        for rel in tree.findall(".//r:Relationship", ns):
            rid = rel.get("Id")
            target = rel.get("Target")
            mode = rel.get("TargetMode")
            if rid and target and mode == "External":
                rels[rid] = target
    except Exception:
        pass
    return rels


def _parse_xml_text(xml_content: bytes, rels: Dict[str, str]) -> str:
    """Extracts text and hyperlinks from XML content."""
    try:
        root = ET.fromstring(xml_content)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
              "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        
        text_parts = []
        
        # Iterate over paragraphs
        for p in root.findall(".//w:p", ns):
            p_text = []
            
            # Helper to extract text from a node's children in order
            # Since ElementTree findall doesn't give mixed children easily,
            # we iterate over all children.
            for child in p:
                tag = child.tag
                if tag.endswith("hyperlink"):
                    rid = child.get(f"{{{ns['r']}}}id")
                    url = rels.get(rid)
                    
                    # Get text inside hyperlink
                    link_texts = [t.text or "" for t in child.findall(".//w:t", ns)]
                    link_text = "".join(link_texts)
                    
                    if url:
                        p_text.append(f"{link_text} ({url})")
                    else:
                        p_text.append(link_text)
                elif tag.endswith("r"): # Run
                    texts = [t.text or "" for t in child.findall(".//w:t", ns)]
                    p_text.append("".join(texts))
            
            # Also handle simple text directly in p (rare but possible in some xml variants?)
            # Usually text is in w:r/w:t.
            
            if p_text:
                text_parts.append("".join(p_text))
                
        return "\n".join(text_parts)
    except Exception:
        return ""


def extract_text_from_docx_bytes(data: bytes) -> str:
    bio = io.BytesIO(data)
    text_content = []
    
    try:
        with zipfile.ZipFile(bio) as zf:
            # 1. Relationships for main document
            doc_rels = _get_relationships(zf, "word/_rels/document.xml.rels")
            
            # 2. Main content
            if "word/document.xml" in zf.namelist():
                with zf.open("word/document.xml") as f:
                    text_content.append(_parse_xml_text(f.read(), doc_rels))
            
            # 3. Headers and Footers
            for name in zf.namelist():
                if name.startswith("word/header") or name.startswith("word/footer"):
                    if name.endswith(".xml"):
                        # rels file for header1.xml is word/_rels/header1.xml.rels
                        # Filename logic: word/header1.xml -> word/_rels/header1.xml.rels
                        base_name = name.split("/")[-1]
                        rel_name = f"word/_rels/{base_name}.rels"
                        rels = _get_relationships(zf, rel_name)
                        
                        with zf.open(name) as f:
                            content = _parse_xml_text(f.read(), rels)
                            if content:
                                # Append with a separator to indicate it's a separate section
                                text_content.append(f"\n--- {base_name} ---\n{content}")
    except zipfile.BadZipFile:
        return ""
    except Exception:
        return ""

    return "\n\n".join(text_content)


def docx_to_pdf_bytes(data: bytes) -> bytes:
    text = extract_text_from_docx_bytes(data)
    return make_simple_pdf_from_text(text)


def _make_simple_docx_from_text(text: str) -> bytes:
    """Fallback: build a basic DOCX with plain paragraphs."""
    doc = Document()
    for line in text.splitlines():
        doc.add_paragraph(line)
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def pdf_bytes_to_docx_bytes(pdf_bytes: bytes) -> Optional[bytes]:
    """
    Converts PDF bytes to DOCX bytes using pdf2docx with maximum fidelity.
    Falls back to a simple text DOCX if conversion fails.
    """
    if not isinstance(pdf_bytes, (bytes, bytearray)) or len(pdf_bytes) < 4:
        return None
    tmp_pdf = None
    tmp_docx = None
    try:
        tmp_pdf = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")
        tmp_docx = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        try:
            tmp_pdf.write(pdf_bytes)
            tmp_pdf.flush()
        finally:
            tmp_pdf.close()
            tmp_docx.close()
        try:
            logger.info("pdf2docx conversion start")
            cv = Converter(tmp_pdf.name)
            cv.convert(tmp_docx.name, start=0, end=None)
            cv.close()
            logger.info("pdf2docx conversion complete")
        except Exception:
            logger.exception("pdf2docx conversion failed")
            simple = _make_simple_docx_from_text("PDF to DOCX conversion failed; fallback document.")
            return simple
        with open(tmp_docx.name, "rb") as f:
            data = f.read()
        if is_docx_bytes(data):
            return data
        logger.error("pdf2docx produced invalid DOCX bytes")
        return _make_simple_docx_from_text("PDF to DOCX conversion produced invalid DOCX; fallback document.")
    except Exception:
        logger.exception("PDF→DOCX pipeline failed")
        return None
    finally:
        for path in (getattr(tmp_pdf, "name", None), getattr(tmp_docx, "name", None)):
            if path and os.path.exists(path):
                try:
                    os.remove(path)
                except Exception:
                    pass


def apply_font_to_docx(docx_bytes: bytes, font_name: str) -> Optional[bytes]:
    if not docx_bytes or not font_name:
        return docx_bytes
    try:
        bio = io.BytesIO(docx_bytes)
        doc = Document(bio)
        try:
            style = doc.styles["Normal"]
            style.font.name = font_name
            try:
                style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
            except Exception:
                pass
        except Exception:
            pass
        for p in doc.paragraphs:
            for r in p.runs:
                r.font.name = font_name
                try:
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                except Exception:
                    pass
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for r in p.runs:
                            r.font.name = font_name
                            try:
                                r._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
                            except Exception:
                                pass
        out = io.BytesIO()
        doc.save(out)
        data = out.getvalue()
        if is_docx_bytes(data):
            return data
        return docx_bytes
    except Exception:
        logger.exception("apply_font_to_docx failed")
        return docx_bytes


def enforce_white_background_docx(docx_bytes: bytes) -> Optional[bytes]:
    """
    Remove shading/highlight on runs and cells, ensure white background.
    """
    if not docx_bytes:
        return docx_bytes
    try:
        bio = io.BytesIO(docx_bytes)
        doc = Document(bio)
        ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        def _clear_run_bg(run):
            try:
                rPr = run._element.rPr
                if rPr is not None:
                    # Remove shading
                    shd = rPr.find(qn("w:shd"))
                    while shd is not None:
                        rPr.remove(shd)
                        shd = rPr.find(qn("w:shd"))
                    # Remove highlight
                    highlight = rPr.find(qn("w:highlight"))
                    if highlight is not None:
                        rPr.remove(highlight)
            except Exception:
                pass

        def _clear_para_bg(p):
            for r in p.runs:
                _clear_run_bg(r)

        # Clear paragraphs
        for p in doc.paragraphs:
            _clear_para_bg(p)

        # Clear tables (cells + their paragraphs)
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    try:
                        tcPr = cell._tc.get_or_add_tcPr()
                        shd = tcPr.find(qn("w:shd"))
                        while shd is not None:
                            tcPr.remove(shd)
                            shd = tcPr.find(qn("w:shd"))
                    except Exception:
                        pass
                    for p in cell.paragraphs:
                        _clear_para_bg(p)

        # Ensure no document-level background
        try:
            bg_elems = doc.element.xpath("//w:background", namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"})
            for bg in bg_elems:
                parent = bg.getparent()
                if parent is not None:
                    parent.remove(bg)
        except Exception:
            pass

        out = io.BytesIO()
        doc.save(out)
        data = out.getvalue()
        if is_docx_bytes(data):
            return data
        return docx_bytes
    except Exception:
        logger.exception("enforce_white_background_docx failed")
        return docx_bytes


def set_docx_vertical_margins_px(docx_bytes: bytes, top_px: int = 5, bottom_px: int = 5) -> Optional[bytes]:
    if not docx_bytes:
        return docx_bytes
    try:
        bio = io.BytesIO(docx_bytes)
        doc = Document(bio)
        top_pt = Pt(top_px * 0.75)
        bottom_pt = Pt(bottom_px * 0.75)
        for sec in doc.sections:
            sec.top_margin = top_pt
            sec.bottom_margin = bottom_pt
        out = io.BytesIO()
        doc.save(out)
        data = out.getvalue()
        if is_docx_bytes(data):
            return data
        return docx_bytes
    except Exception:
        logger.exception("set_docx_vertical_margins_px failed")
        return docx_bytes

def docx_bytes_to_jpeg_pages(docx_bytes: bytes, dpi: int = 144, quality: int = 85) -> List[bytes]:
    if not isinstance(docx_bytes, (bytes, bytearray)) or not is_docx_bytes(docx_bytes):
        raise ValueError("Input must be valid DOCX bytes")
    if dpi <= 0:
        raise ValueError("dpi must be positive")
    if not (1 <= quality <= 95):
        raise ValueError("quality must be between 1 and 95")

    pdf_bytes = docx_to_pdf_bytes(docx_bytes)
    if not pdf_bytes:
        raise ValueError("DOCX to PDF conversion failed")

    return pdf_bytes_to_jpeg_pages(pdf_bytes, dpi=dpi, quality=quality)

